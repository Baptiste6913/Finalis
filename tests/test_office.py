"""PowerPoint and Word exports of the corrected material: built in the page, opened with python-pptx /
python-docx, converted by LibreOffice. Usage: python3 tests/test_office.py"""
import asyncio, sys, os, base64, subprocess
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from harness import *
FAKE = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'test_flow.py')).read().split('FAKE = r"""')[1].split('"""')[0]
OUT = os.path.join(ROOT, 'tests', 'out'); os.makedirs(OUT, exist_ok=True)
async def main():
    ok = True
    async with async_playwright() as pw:
        browser, page = await open_page(pw, width=1280, height=1000, init_script=FAKE)
        await page.wait_for_timeout(500)
        await page.click('#sample-slot button'); await page.wait_for_selector('#filechip:not([hidden])', timeout=30000)
        await page.click('#btn-submit-form'); await page.wait_for_selector('#btn-start'); await page.click('#btn-start')
        await page.wait_for_function("document.getElementById('wh-status-text').textContent.includes('complete')", timeout=60000)
        await page.wait_for_timeout(400)
        await page.evaluate("UI.applyAllFixes()"); await page.wait_for_timeout(200)
        for kind in ('pptx', 'docx'):
            b64 = await page.evaluate("""async (kind) => { const base = UI.S.doc; const changes = Fix.changeLog(UI.S.edits, UI.S.result.findings); const pts = UI.S.result.findings; const b = kind === 'pptx' ? await Exports.pptx(base, changes, pts, { title: 'Quartus' }) : await Exports.docx(base, changes, pts, { title: 'Quartus', version: 1 }); let s=''; for (let i=0;i<b.length;i+=1) s+=String.fromCharCode(b[i]); return btoa(s); }""", kind)
            out = os.path.join(OUT, 'quartus-corrected.' + kind); open(out, 'wb').write(base64.b64decode(b64))
            print(kind, os.path.getsize(out), 'bytes')
        await browser.close()
    from pptx import Presentation
    prs = Presentation(os.path.join(OUT, 'quartus-corrected.pptx'))
    shapes = [len(s.shapes) for s in prs.slides]
    texts = [sh.text_frame.text[:60] for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    print('pptx slides', len(prs.slides), 'shapes per slide', shapes[:6], 'slide 1 texts', texts)
    ok = ok and len(prs.slides) == 29 and any('Institutional' in t for t in texts)
    from docx import Document
    d = Document(os.path.join(OUT, 'quartus-corrected.docx'))
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    print('docx paragraphs', len(paras), 'images', len(d.inline_shapes), 'tables', len(d.tables), 'first', paras[:3])
    ok = ok and len(d.inline_shapes) >= 3 and len(d.tables) == 1
    for kind in ('pptx', 'docx'):
        r = subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', OUT, os.path.join(OUT, 'quartus-corrected.' + kind)], capture_output=True, text=True, timeout=240)
        pdf = os.path.join(OUT, 'quartus-corrected.pdf')
        good = r.returncode == 0 and os.path.exists(pdf)
        if good:
            pages = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
            print(kind, 'LibreOffice ->', [l for l in pages.splitlines() if l.startswith('Pages')])
            os.rename(pdf, os.path.join(OUT, 'quartus-corrected-' + kind + '.pdf'))
        else: print(kind, 'LibreOffice FAILED', r.stderr[-300:])
        ok = ok and good
    print('ALL OK' if ok else 'FAILURES'); sys.exit(0 if ok else 1)
asyncio.run(main())
